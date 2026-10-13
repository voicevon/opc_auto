"""
OPC Auto 电子税务局进项发票底账手动导入服务 (过渡方案)
在税局自动化同步落地前，支持将电子税务局导出的 Excel(xlsx) / CSV 底账批量导入 tax_inbound_invoices。
"""

import base64
import csv
import io
import os
from datetime import datetime, date
from .database import get_db_connection

# 底账表头字段映射 (单一真理源)：目标字段 -> 识别关键词 (按优先级排列，命中即止)
HEADER_KEYWORDS = {
    "invoice_number":     ["发票号码"],
    "invoice_code":       ["发票代码"],
    "issue_date":         ["开票日期"],
    "seller_tax_id":      ["销方税号", "销售方税号", "销方纳税人识别号", "纳税人识别号"],
    "seller_name":        ["销方名称", "销售方名称", "销方单位名称", "销售方"],
    "amount_without_tax": ["不含税金额", "金额"],
    "tax_amount":         ["税额"],
    "total_amount":       ["价税合计", "价税合计金额", "合计金额"],
    "invoice_type":       ["发票类型", "票种"],
    "invoice_state":      ["发票状态", "状态"],
}

# 必填字段
REQUIRED_FIELDS = ["invoice_number", "issue_date", "seller_name", "total_amount"]


def _normalize_header(name) -> str:
    """表头归一化：去空白与括号差异"""
    if name is None:
        return ""
    s = str(name).strip()
    for ch in (" ", "\u3000", "(", ")", "（", "）"):
        s = s.replace(ch, "")
    return s


def _build_header_index(headers):
    """将表头行映射为 {目标字段: 列索引}，关键词优先级高的先匹配"""
    normalized = [_normalize_header(h) for h in headers]
    mapping = {}
    for field, keywords in HEADER_KEYWORDS.items():
        for kw in keywords:
            kw_norm = _normalize_header(kw)
            for idx, h in enumerate(normalized):
                if h and kw_norm in h and idx not in mapping.values():
                    mapping[field] = idx
                    break
            if field in mapping:
                break
    return mapping


def _to_float(val):
    if val is None or str(val).strip() == "":
        return None
    try:
        return float(str(val).replace("￥", "").replace("¥", "").replace(",", "").replace(" ", ""))
    except (ValueError, TypeError):
        return None


def _to_date_str(val):
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return s


def _normalize_state(val):
    s = str(val or "").strip()
    if "红" in s:
        return "red"
    if "作废" in s:
        return "cancelled"
    return "normal"


def _parse_rows(filename: str, raw_bytes: bytes):
    """根据文件扩展名解析为行记录列表 (每行为 list)"""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    if ext == ".csv":
        for codec in ("utf-8-sig", "gbk"):
            try:
                text = raw_bytes.decode(codec)
                break
            except (UnicodeDecodeError, LookupError):
                continue
        else:
            text = raw_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        return [row for row in reader if any(str(c).strip() for c in row)]
    raise ValueError(f"不支持的文件格式: {ext} (仅支持 .xlsx / .csv，税局底账可另存为 xlsx 或 CSV)")


def import_invoice_file(filename: str, content_base64: str):
    """解析底账文件并导入 tax_inbound_invoices (raw_source=manual_import)"""
    try:
        raw_bytes = base64.b64decode(content_base64)
        rows = _parse_rows(filename, raw_bytes)
    except Exception as e:
        return {"success": False, "message": f"文件解析失败: {e}", "imported": 0, "skipped": 0}

    if len(rows) < 2:
        return {"success": False, "message": "文件为空或缺少数据行", "imported": 0, "skipped": len(rows)}

    header_idx = _build_header_index(rows[0])
    missing_required = [f for f in REQUIRED_FIELDS if f not in header_idx]
    if missing_required:
        return {
            "success": False,
            "message": f"表头缺少必要列: {missing_required}，请使用电子税务局底账导出原表",
            "detected_headers": [str(h) for h in rows[0] if h],
            "imported": 0, "skipped": len(rows) - 1
        }

    conn = get_db_connection()
    cursor = conn.cursor()
    imported, skipped, errors = 0, 0, []

    for row_no, row in enumerate(rows[1:], start=2):
        def cell(field):
            idx = header_idx.get(field)
            return row[idx] if idx is not None and idx < len(row) else None

        invoice_number = str(cell("invoice_number") or "").strip()
        total_amount = _to_float(cell("total_amount"))
        if not invoice_number or total_amount is None:
            skipped += 1
            continue

        amount_without_tax = _to_float(cell("amount_without_tax"))
        tax_amount = _to_float(cell("tax_amount"))
        if amount_without_tax is None:
            amount_without_tax = round(total_amount - (tax_amount or 0.0), 2)
        if tax_amount is None:
            tax_amount = round(total_amount - amount_without_tax, 2)

        try:
            cursor.execute("""
            INSERT OR IGNORE INTO tax_inbound_invoices (
                invoice_code, invoice_number, issue_date, seller_tax_id, seller_name,
                amount_without_tax, tax_amount, total_amount, invoice_type, invoice_state, raw_source
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'manual_import')
            """, (
                str(cell("invoice_code") or "").strip(),
                invoice_number,
                _to_date_str(cell("issue_date")),
                str(cell("seller_tax_id") or "").strip(),
                str(cell("seller_name") or "").strip(),
                amount_without_tax,
                tax_amount,
                total_amount,
                str(cell("invoice_type") or "数电普票").strip(),
                _normalize_state(cell("invoice_state"))
            ))
            if cursor.rowcount > 0:
                imported += 1
            else:
                skipped += 1  # 票号已存在 (UNIQUE 幂等去重)
        except Exception as e:
            errors.append(f"第{row_no}行: {e}")

    conn.commit()
    conn.close()

    return {
        "success": True,
        "message": f"导入完成：新增 {imported} 张，跳过 {skipped} 张" + (f"，异常 {len(errors)} 行" if errors else ""),
        "imported": imported,
        "skipped": skipped,
        "errors": errors[:5]
    }
