"""
OPC Auto 淘宝订单官方导出文件导入服务
消费淘宝「已买到的宝贝 → 导出订单」生成的 xlsx，作为订单主数据（实付金额/数量/运费/下单时间）的权威数据源。
CDP 拟人化采集仅负责发票状态与开票动作；金额等主数据以本导入通道覆盖修正。
"""

import base64
import io
from datetime import datetime, date
from .database import get_db_connection
from .reconciliation_engine import run_reconciliation

# 淘宝导出表头字段映射 (单一真理源)：目标字段 -> 识别关键词 (按优先级排列，命中即止)
HEADER_KEYWORDS = {
    "order_no":     ["订单号"],
    "order_time":   ["订单提交时间"],
    "order_status": ["订单状态"],
    "seller_name":  ["店铺名称"],
    "item_name":    ["商品名称"],
    "item_model":   ["型号款式"],
    "quantity":     ["商品数量"],
    "goods_amount": ["商品金额"],
    "pay_amount":   ["实付金额"],
    "freight":      ["运费"],
}

# 必填字段
REQUIRED_FIELDS = ["order_no", "order_time", "pay_amount"]


def _normalize_header(name) -> str:
    if name is None:
        return ""
    s = str(name).strip()
    for ch in (" ", "\u3000", "(", ")", "（", "）"):
        s = s.replace(ch, "")
    return s


def _build_header_index(headers):
    normalized = [_normalize_header(h) for h in headers]
    mapping = {}
    for field, keywords in HEADER_KEYWORDS.items():
        for kw in keywords:
            kw_norm = _normalize_header(kw)
            for idx, h in enumerate(normalized):
                if h and kw_norm in h and idx not in mapping.values():
                    mapping[field] = idx
                    break
            if field in mapping:
                break
    return mapping


def _to_float(val):
    if val is None or str(val).strip() == "":
        return None
    try:
        return float(str(val).replace("￥", "").replace("¥", "").replace(",", "").replace(" ", ""))
    except (ValueError, TypeError):
        return None


def _to_datetime_str(val):
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    return str(val).strip()


def _parse_rows(filename: str, raw_bytes: bytes):
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    raise ValueError("仅支持淘宝「导出订单」生成的 .xlsx 文件")


def import_order_export_file(filename: str, content_base64: str):
    """解析淘宝订单导出文件并 upsert 入 procurement_orders (实付金额为权威值)"""
    try:
        raw_bytes = base64.b64decode(content_base64)
        rows = _parse_rows(filename, raw_bytes)
    except Exception as e:
        return {"success": False, "message": f"文件解析失败: {e}", "imported": 0, "skipped": 0}

    if len(rows) < 2:
        return {"success": False, "message": "文件为空或缺少数据行", "imported": 0, "skipped": len(rows)}

    header_idx = _build_header_index(rows[0])
    missing_required = [f for f in REQUIRED_FIELDS if f not in header_idx]
    if missing_required:
        return {
            "success": False,
            "message": f"表头缺少必要列: {missing_required}，请使用淘宝「导出订单」原表",
            "detected_headers": [str(h) for h in rows[0] if h],
            "imported": 0, "skipped": len(rows) - 1
        }

    conn = get_db_connection()
    cursor = conn.cursor()
    imported, skipped, errors = 0, 0, []
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for row_no, row in enumerate(rows[1:], start=2):
        def cell(field):
            idx = header_idx.get(field)
            return row[idx] if idx is not None and idx < len(row) else None

        order_no = str(cell("order_no") or "").strip()
        pay_amount = _to_float(cell("pay_amount"))
        order_time = _to_datetime_str(cell("order_time"))
        if not order_no or pay_amount is None:
            skipped += 1
            continue

        seller = str(cell("seller_name") or "").strip()
        item_name = str(cell("item_name") or "").strip()
        item_model = str(cell("item_model") or "").strip()
        item_desc = f"{item_name}（{item_model}）" if item_model else item_name
        quantity = _to_float(cell("quantity"))
        goods_amount = _to_float(cell("goods_amount"))
        freight = _to_float(cell("freight"))
        order_status = str(cell("order_status") or "").strip()

        meta_parts = [f"淘宝导出导入@{now_str}", f"订单状态: {order_status}" if order_status else ""]
        if quantity is not None:
            meta_parts.append(f"数量x{int(quantity) if quantity == int(quantity) else quantity}")
        if goods_amount is not None:
            meta_parts.append(f"商品额{goods_amount}")
        if freight:
            meta_parts.append(f"运费{freight}")

        try:
            cursor.execute("""
            INSERT INTO procurement_orders (
                channel, channel_name, order_no, order_time, item_desc, amount, seller_name, invoice_status, notes
            ) VALUES ('taobao', '淘宝企业购', ?, ?, ?, ?, ?, 'unapplied', ?)
            ON CONFLICT(order_no) DO UPDATE SET
                order_time=excluded.order_time,
                amount=excluded.amount,
                seller_name=excluded.seller_name,
                item_desc=excluded.item_desc,
                notes=excluded.notes
            """, (order_no, order_time, item_desc, pay_amount, seller, " | ".join(p for p in meta_parts if p)))
            if cursor.rowcount > 0:
                imported += 1
            else:
                skipped += 1
        except Exception as e:
            errors.append(f"第{row_no}行: {e}")

    conn.commit()
    conn.close()

    # 导入后重跑对账 (实付金额修正可能改变配对结果)
    run_reconciliation()

    return {
        "success": True,
        "message": f"导入完成：更新/新增 {imported} 笔，跳过 {skipped} 笔" + (f"，异常 {len(errors)} 行" if errors else ""),
        "imported": imported,
        "skipped": skipped,
        "errors": errors[:5]
    }
